"""Upload parsing: CSV / JSON / XLSX / TXT to row dicts."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from apps.api.core.errors import DatasetValidationError

ALLOWED_CONTENT_TYPES: frozenset[str] = frozenset(
    {
        "text/csv",
        "application/csv",
        "application/json",
        "text/plain",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }
)


def parse_upload(filename: str, data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Parse an uploaded file into ``(rows, columns)``.

    Supports .csv, .json, .jsonl, .txt, .xlsx (lazy optional openpyxl).
    """
    if not filename:
        raise DatasetValidationError("Upload missing filename")
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _parse_csv(data)
    if lower.endswith(".jsonl"):
        return _parse_jsonl(data)
    if lower.endswith(".json"):
        return _parse_json(data)
    if lower.endswith(".txt"):
        return _parse_txt(data)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(data)
    raise DatasetValidationError(
        f"Unsupported file type: {filename}",
        {"filename": filename},
    )


def _parse_csv(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = [dict(row) for row in reader]
    columns = list(reader.fieldnames or [])
    if not columns:
        raise DatasetValidationError("CSV has no header row")
    return rows, columns


def _parse_json(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        payload = json.loads(data.decode("utf-8"))
    except json.JSONDecodeError as exc:
        raise DatasetValidationError(f"Invalid JSON: {exc}") from exc
    if isinstance(payload, dict) and "data" in payload and isinstance(payload["data"], list):
        payload = payload["data"]
    if not isinstance(payload, list):
        raise DatasetValidationError("JSON payload must be a list of objects")
    rows = [r for r in payload if isinstance(r, dict)]
    columns = sorted({k for r in rows for k in r})
    return rows, columns


def _parse_jsonl(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    rows: list[dict[str, Any]] = []
    for raw_line in data.decode("utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        try:
            parsed = json.loads(line)
        except json.JSONDecodeError as exc:
            raise DatasetValidationError(f"Invalid JSONL line: {exc}") from exc
        if isinstance(parsed, dict):
            rows.append(parsed)
    columns = sorted({k for r in rows for k in r})
    return rows, columns


def _parse_txt(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    text = data.decode("utf-8", errors="replace")
    rows = [{"text": line.strip()} for line in text.splitlines() if line.strip()]
    return rows, ["text"]


def _parse_xlsx(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DatasetValidationError(
            "openpyxl is not installed. Install the 'ml' extra to enable XLSX uploads.",
        ) from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise DatasetValidationError("XLSX has no active sheet")
    iterator = ws.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise DatasetValidationError("XLSX is empty") from exc
    columns = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header_row)]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        row: dict[str, Any] = {col: values[i] for i, col in enumerate(columns) if i < len(values)}
        rows.append(row)
    return rows, columns
